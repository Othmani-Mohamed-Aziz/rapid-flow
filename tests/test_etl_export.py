from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from app.config.study_schema_provider import load_study_schema_from_json
from app.etl.ecrf_export_policy import partition_cell_updates_for_ecrf
from app.etl.export import EcrfExportService
from app.etl.overwrite_policy import EcrfTemplateSnapshot, XlsOverwritePolicy
from app.etl.xls_export import XlsEcrfExporter
from app.etl.xls_utils import build_header_map, is_ecrf_cell_empty
from app.extraction.planner import plan_extraction_jobs
from app.schemas.enums import DocumentType, FieldFamily
from app.schemas.models import EcrfCellUpdate, ExtractedObservation, FieldCandidate, PipelineResult

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SCHEMA = PROJECT_ROOT / "data" / "study_schema_default.json"


def _make_obs(value: int | str) -> ExtractedObservation:
    return ExtractedObservation(
        observation_id="o1",
        field_family=FieldFamily.HEMATOLOGY,
        normalized_value=value,
        confidence=0.9,
        evidence_text=str(value),
        extraction_method="test",
    )


def _make_update(
    *,
    patient_id: str,
    column_key: str,
    value: int | str | bool,
) -> EcrfCellUpdate:
    obs = _make_obs(value if isinstance(value, (int, str)) else value)
    cand = FieldCandidate(
        field_name=column_key,
        observation=obs,
        target_column=column_key,
        confidence=0.9,
    )
    return EcrfCellUpdate(
        patient_id=patient_id,
        study_id="EXAMPLE",
        column_key=column_key,
        value=value,
        provenance=cand,
        validated=True,
    )


def _make_result(
    *,
    patient_id: str,
    updates: list[EcrfCellUpdate],
    ma_patient_key: str | None = None,
) -> PipelineResult:
    return PipelineResult(
        document_id="d1",
        patient_id=patient_id,
        study_id="EXAMPLE",
        document_type=DocumentType.LAB_BLOOD_PANEL,
        observations=[u.provenance.observation for u in updates],
        cell_updates=updates,
        audit_trail=[],
        ma_patient_key=ma_patient_key or patient_id,
    )


def make_minimal_ma_template(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Global_CC"
    headers = [
        "ID_current_base",
        "Center",
        "PLT_start_AtezoBev_D0",
        "AST_start_AtezoBev_D0",
        "Response_at_the_first_imaging_RECIST",
    ]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=2, column=col, value=header)
    ws.cell(row=3, column=1, value=220)
    ws.cell(row=3, column=2, value="Beaujon")
    ws.cell(row=3, column=3, value=120)
    wb.save(path)
    wb.close()


def test_is_ecrf_cell_empty() -> None:
    assert is_ecrf_cell_empty(None) is True
    assert is_ecrf_cell_empty("") is True
    assert is_ecrf_cell_empty("   ") is True
    assert is_ecrf_cell_empty(0) is False
    assert is_ecrf_cell_empty("NA") is True
    assert is_ecrf_cell_empty("42") is False
    assert is_ecrf_cell_empty("x") is False


def test_snapshot_detects_filled_columns(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    make_minimal_ma_template(template)
    snap = EcrfTemplateSnapshot(template_path=template, ma_patient_key="220")
    assert snap.is_column_filled("PLT_start_AtezoBev_D0") is True
    assert snap.is_column_filled("AST_start_AtezoBev_D0") is False
    filled = snap.filled_target_columns({"PLT_start_AtezoBev_D0", "AST_start_AtezoBev_D0"})
    assert filled == {"PLT_start_AtezoBev_D0"}


def test_planner_excludes_filled_columns() -> None:
    schema = load_study_schema_from_json(DEFAULT_SCHEMA)
    all_jobs = plan_extraction_jobs(schema, DocumentType.LAB_BLOOD_PANEL)
    filtered = plan_extraction_jobs(
        schema,
        DocumentType.LAB_BLOOD_PANEL,
        exclude_target_columns={"PLT_start_AtezoBev_D0"},
    )
    all_fields = {name for job in all_jobs for name in job.field_names}
    filtered_fields = {name for job in filtered for name in job.field_names}
    assert "PLT_start_AtezoBev_D0" in all_fields
    assert "PLT_start_AtezoBev_D0" not in filtered_fields
    assert len(filtered) <= len(all_jobs)


def test_export_writes_json_and_csv(tmp_path: Path) -> None:
    upd = _make_update(patient_id="p1", column_key="PLT_start_AtezoBev_D0", value=150)
    result = _make_result(patient_id="p1", updates=[upd])
    template = tmp_path / "template.xlsx"
    make_minimal_ma_template(template)
    svc = EcrfExportService(
        template_xlsx_path=template,
        overwrite_policy=XlsOverwritePolicy.ALWAYS,
    )
    paths = svc.export(result, tmp_path / "out", basename="unit")
    assert Path(paths["json"]).is_file()
    assert Path(paths["csv"]).is_file()
    text = Path(paths["csv"]).read_text(encoding="utf-8")
    assert "PLT_start_AtezoBev_D0" in text


def test_xls_export_always_overwrites_filled_cell(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    make_minimal_ma_template(template)
    dest = tmp_path / "out.xlsx"
    upd = _make_update(patient_id="220", column_key="PLT_start_AtezoBev_D0", value=155)
    result = _make_result(patient_id="220", updates=[upd])

    exporter = XlsEcrfExporter(
        template_path=template,
        overwrite_policy=XlsOverwritePolicy.ALWAYS,
    )
    exporter.export_patient_row(dest_path=dest, result=result)

    wb = openpyxl.load_workbook(dest, data_only=True)
    ws = wb["Global_CC"]
    assert ws.cell(row=3, column=3).value == 155
    wb.close()


def test_xls_export_empty_only_preserves_filled_cell(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    make_minimal_ma_template(template)
    dest = tmp_path / "out.xlsx"
    upd = _make_update(patient_id="220", column_key="PLT_start_AtezoBev_D0", value=155)
    result = _make_result(patient_id="220", updates=[upd])

    stats = XlsEcrfExporter(
        template_path=template,
        overwrite_policy=XlsOverwritePolicy.EMPTY_ONLY,
    ).merge_patient_row_with_stats(workbook_path=template, result=result)
    assert stats.written == 0
    assert stats.skipped_filled == 1

    XlsEcrfExporter(
        template_path=template,
        overwrite_policy=XlsOverwritePolicy.EMPTY_ONLY,
    ).export_patient_row(dest_path=dest, result=result)

    wb = openpyxl.load_workbook(dest, data_only=True)
    ws = wb["Global_CC"]
    assert ws.cell(row=3, column=3).value == 120
    wb.close()


def test_xls_export_empty_only_writes_empty_column(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    make_minimal_ma_template(template)
    dest = tmp_path / "out.xlsx"
    upd = _make_update(patient_id="220", column_key="AST_start_AtezoBev_D0", value=42)
    result = _make_result(patient_id="220", updates=[upd])

    XlsEcrfExporter(
        template_path=template,
        overwrite_policy=XlsOverwritePolicy.EMPTY_ONLY,
    ).export_patient_row(dest_path=dest, result=result)

    wb = openpyxl.load_workbook(dest, data_only=True)
    assert wb["Global_CC"].cell(row=3, column=4).value == 42
    wb.close()


def test_xls_export_never_writes(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    make_minimal_ma_template(template)
    dest = tmp_path / "out.xlsx"
    upd = _make_update(patient_id="220", column_key="AST_start_AtezoBev_D0", value=42)
    result = _make_result(patient_id="220", updates=[upd])

    stats = XlsEcrfExporter(
        template_path=template,
        overwrite_policy=XlsOverwritePolicy.NEVER,
    ).merge_patient_row_with_stats(workbook_path=template, result=result)
    assert stats.skipped_policy_never == 1
    assert stats.written == 0

    XlsEcrfExporter(
        template_path=template,
        overwrite_policy=XlsOverwritePolicy.NEVER,
    ).export_patient_row(dest_path=dest, result=result)

    wb = openpyxl.load_workbook(dest, data_only=True)
    assert wb["Global_CC"].cell(row=3, column=4).value is None
    wb.close()


def test_xls_export_appends_unknown_patient(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    make_minimal_ma_template(template)
    dest = tmp_path / "out.xlsx"
    upd = _make_update(patient_id="999", column_key="AST_start_AtezoBev_D0", value=42)
    result = _make_result(patient_id="999", updates=[upd])

    XlsEcrfExporter(
        template_path=template,
        overwrite_policy=XlsOverwritePolicy.ALWAYS,
    ).export_patient_row(dest_path=dest, result=result)

    wb = openpyxl.load_workbook(dest, data_only=True)
    ws = wb["Global_CC"]
    assert ws.cell(row=4, column=1).value == 999
    assert ws.cell(row=4, column=4).value == 42
    wb.close()


def test_xls_export_column_alias_recist(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    make_minimal_ma_template(template)
    dest = tmp_path / "out.xlsx"
    upd = _make_update(
        patient_id="220",
        column_key="Response_at_first_imaging_RECIST",
        value="SD",
    )
    result = _make_result(patient_id="220", updates=[upd])

    XlsEcrfExporter(
        template_path=template,
        overwrite_policy=XlsOverwritePolicy.ALWAYS,
    ).export_patient_row(dest_path=dest, result=result)

    wb = openpyxl.load_workbook(dest, data_only=True)
    ws = wb["Global_CC"]
    assert ws.cell(row=3, column=5).value == "SD"
    wb.close()


def test_export_service_includes_xlsx_path(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    make_minimal_ma_template(template)
    upd = _make_update(patient_id="220", column_key="AST_start_AtezoBev_D0", value=88)
    result = _make_result(patient_id="220", updates=[upd])
    svc = EcrfExportService(
        template_xlsx_path=template,
        overwrite_policy=XlsOverwritePolicy.ALWAYS,
    )
    paths = svc.export(result, tmp_path / "out")
    assert "xlsx" in paths
    assert Path(paths["xlsx"]).is_file()


def test_export_skips_xlsx_when_template_missing(tmp_path: Path) -> None:
    upd = _make_update(patient_id="p1", column_key="PLT_start_AtezoBev_D0", value=1)
    result = _make_result(patient_id="p1", updates=[upd])
    svc = EcrfExportService(template_xlsx_path=tmp_path / "missing.xlsx")
    paths = svc.export(result, tmp_path / "out")
    assert "xlsx" not in paths


def test_partition_empty_only_suppresses_filled_columns(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    make_minimal_ma_template(template)
    snap = EcrfTemplateSnapshot(template_path=template, ma_patient_key="220")
    filled = _make_update(patient_id="220", column_key="PLT_start_AtezoBev_D0", value=155)
    empty_col = _make_update(patient_id="220", column_key="AST_start_AtezoBev_D0", value=42)
    accepted, suppressed = partition_cell_updates_for_ecrf(
        [filled, empty_col],
        policy=XlsOverwritePolicy.EMPTY_ONLY,
        snapshot=snap,
        column_aliases={},
    )
    assert len(accepted) == 1
    assert accepted[0].column_key == "AST_start_AtezoBev_D0"
    assert any(s["reason"] == "column_already_filled" for s in suppressed)


def test_partition_never_clears_all_updates() -> None:
    upd = _make_update(patient_id="220", column_key="AST_start_AtezoBev_D0", value=42)
    snap = EcrfTemplateSnapshot(template_path=Path("missing.xlsx"), ma_patient_key="220")
    accepted, suppressed = partition_cell_updates_for_ecrf(
        [upd],
        policy=XlsOverwritePolicy.NEVER,
        snapshot=snap,
        column_aliases={},
    )
    assert accepted == []
    assert suppressed[0]["reason"] == "policy_never"


def test_write_to_master_workbook_updates_template_in_place(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    make_minimal_ma_template(template)
    upd = _make_update(patient_id="220", column_key="AST_start_AtezoBev_D0", value=77)
    result = _make_result(patient_id="220", updates=[upd])
    XlsEcrfExporter(
        template_path=template,
        overwrite_policy=XlsOverwritePolicy.ALWAYS,
    ).write_to_master_workbook(result=result)
    wb = openpyxl.load_workbook(template, data_only=True)
    assert wb["Global_CC"].cell(row=3, column=4).value == 77
    wb.close()


def test_build_header_map_records_duplicate_headers() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.cell(row=2, column=1, value="DupCol")
    ws.cell(row=2, column=2, value="DupCol")
    result = build_header_map(ws, 2)
    assert result.duplicate_headers == ["DupCol"]
    assert result.mapping["DupCol"] == 1
    wb.close()


def test_export_json_csv_match_ecrf_cell_updates_empty_only(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    make_minimal_ma_template(template)
    snap = EcrfTemplateSnapshot(template_path=template, ma_patient_key="220")
    filled = _make_update(patient_id="220", column_key="PLT_start_AtezoBev_D0", value=155)
    empty_col = _make_update(patient_id="220", column_key="AST_start_AtezoBev_D0", value=42)
    ecrf_updates, _ = partition_cell_updates_for_ecrf(
        [filled, empty_col],
        policy=XlsOverwritePolicy.EMPTY_ONLY,
        snapshot=snap,
        column_aliases={},
    )
    result = _make_result(patient_id="220", updates=ecrf_updates)
    svc = EcrfExportService(
        template_xlsx_path=template,
        overwrite_policy=XlsOverwritePolicy.EMPTY_ONLY,
        export_xls_update_master_workbook=False,
    )
    paths = svc.export(result, tmp_path / "out")
    csv_text = Path(paths["csv"]).read_text(encoding="utf-8")
    assert "AST_start_AtezoBev_D0" in csv_text
    assert "PLT_start_AtezoBev_D0" not in csv_text
    payload = json.loads(Path(paths["json"]).read_text(encoding="utf-8"))
    assert len(payload["cell_updates"]) == 1
    assert payload["cell_updates"][0]["column_key"] == "AST_start_AtezoBev_D0"


@pytest.mark.skipif(
    not Path("data/MA_Base_example.xlsx").is_file(),
    reason="Gabarit MA local absent (data/ gitignoré)",
)
def test_xls_export_against_real_ma_base(tmp_path: Path) -> None:
    template = Path("data/MA_Base_example.xlsx")
    dest = tmp_path / "ma_out.xlsx"
    upd = _make_update(patient_id="220", column_key="PLT_start_AtezoBev_D0", value=999)
    result = _make_result(patient_id="220", updates=[upd])
    XlsEcrfExporter(
        template_path=template,
        overwrite_policy=XlsOverwritePolicy.ALWAYS,
    ).export_patient_row(dest_path=dest, result=result)
    wb = openpyxl.load_workbook(dest, data_only=True)
    headers = list(wb["Global_CC"].iter_rows(min_row=2, max_row=2, values_only=True))[0]
    col = headers.index("PLT_start_AtezoBev_D0") + 1
    assert wb["Global_CC"].cell(row=3, column=col).value == 999
    wb.close()
