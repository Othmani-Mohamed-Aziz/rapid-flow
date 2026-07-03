"""Export eCRF XLSX (gabarit MA) — une ligne par patient."""

from __future__ import annotations

import logging
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from app.etl.column_aliases import resolve_column_aliases
from app.etl.ecrf_export_policy import should_write_cell_value
from app.etl.overwrite_policy import XlsOverwritePolicy
from app.etl.xls_utils import build_header_map, find_patient_row, resolve_column_key
from app.schemas.models import EcrfCellUpdate, PipelineResult

_LOG = logging.getLogger(__name__)


class XlsExportError(Exception):
    """Erreur métier lors de l'écriture XLS eCRF."""


@dataclass
class XlsWriteStats:
    written: int = 0
    skipped_missing_column: int = 0
    skipped_filled: int = 0
    skipped_policy_never: int = 0

    def to_dict(self) -> dict[str, int]:
        return {
            "written": self.written,
            "skipped_missing_column": self.skipped_missing_column,
            "skipped_filled": self.skipped_filled,
            "skipped_policy_never": self.skipped_policy_never,
        }


class XlsEcrfExporter:
    """
    Écrit les ``cell_updates`` dans un classeur eCRF MA.

    Par défaut le service d'export appelle ``write_to_master_workbook`` (source de vérité).
    ``export_patient_row`` copie d'abord le gabarit (tests / miroir optionnel).
    """

    def __init__(
        self,
        *,
        template_path: Path,
        sheet_name: str = "Global_CC",
        header_row: int = 2,
        patient_id_column: str = "ID_current_base",
        column_aliases: dict[str, str] | None = None,
        overwrite_policy: XlsOverwritePolicy | str = XlsOverwritePolicy.EMPTY_ONLY,
        empty_sentinels: frozenset[str] | None = None,
    ) -> None:
        self.template_path = template_path
        self.sheet_name = sheet_name
        self.header_row = header_row
        self.patient_id_column = patient_id_column
        self.column_aliases = (
            column_aliases if column_aliases is not None else resolve_column_aliases()
        )
        self.overwrite_policy = (
            overwrite_policy
            if isinstance(overwrite_policy, XlsOverwritePolicy)
            else XlsOverwritePolicy(str(overwrite_policy))
        )
        self.empty_sentinels = empty_sentinels

    def write_to_master_workbook(self, *, result: PipelineResult) -> Path:
        """Met à jour le fichier gabarit source (source de vérité unique)."""
        if not self.template_path.is_file():
            msg = f"Gabarit eCRF introuvable : {self.template_path}"
            raise XlsExportError(msg)
        return self.merge_patient_row(workbook_path=self.template_path, result=result)

    def export_patient_row(self, *, dest_path: Path, result: PipelineResult) -> Path:
        """Copie le gabarit vers ``dest_path`` puis fusionne (tests / artefact miroir)."""
        if self.overwrite_policy == XlsOverwritePolicy.NEVER:
            _LOG.info(
                "Export XLS copie ignorée (policy=never) : %s",
                dest_path,
            )
            if dest_path.is_file():
                return dest_path
            if not self.template_path.is_file():
                msg = f"Gabarit eCRF introuvable : {self.template_path}"
                raise XlsExportError(msg)
            dest_path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(self.template_path, dest_path)
            return dest_path

        if not self.template_path.is_file():
            msg = f"Gabarit eCRF introuvable : {self.template_path}"
            raise XlsExportError(msg)
        dest_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(self.template_path, dest_path)
        return self.merge_patient_row(workbook_path=dest_path, result=result)

    def merge_patient_row(self, *, workbook_path: Path, result: PipelineResult) -> Path:
        stats = self.merge_patient_row_with_stats(workbook_path=workbook_path, result=result)
        _LOG.info(
            "Export XLS eCRF : patient=%s policy=%s stats=%s → %s",
            result.ma_patient_key or result.patient_id,
            self.overwrite_policy.value,
            stats.to_dict(),
            workbook_path,
        )
        return workbook_path

    def merge_patient_row_with_stats(
        self,
        *,
        workbook_path: Path,
        result: PipelineResult,
    ) -> XlsWriteStats:
        if self.overwrite_policy == XlsOverwritePolicy.NEVER:
            _LOG.info(
                "Écriture XLS ignorée (policy=never) pour patient %s",
                result.ma_patient_key or result.patient_id,
            )
            return XlsWriteStats(skipped_policy_never=len(result.cell_updates))

        try:
            import openpyxl
        except ImportError as e:
            msg = "openpyxl requis : pip install openpyxl"
            raise XlsExportError(msg) from e

        if not workbook_path.is_file():
            msg = f"Classeur eCRF introuvable : {workbook_path}"
            raise XlsExportError(msg)

        ma_key = result.ma_patient_key or result.patient_id
        wb = openpyxl.load_workbook(workbook_path)
        try:
            if self.sheet_name not in wb.sheetnames:
                msg = f"Feuille {self.sheet_name!r} absente du gabarit (feuilles : {wb.sheetnames})"
                raise XlsExportError(msg)
            ws = wb[self.sheet_name]
            header_result = build_header_map(ws, self.header_row)
            header_map = header_result.mapping
            if header_result.duplicate_headers:
                _LOG.warning(
                    "Colonnes dupliquées ignorées à l'écriture : %s",
                    header_result.duplicate_headers,
                )
            if self.patient_id_column not in header_map:
                msg = (
                    f"Colonne patient {self.patient_id_column!r} absente "
                    f"(ligne {self.header_row})"
                )
                raise XlsExportError(msg)

            patient_col = header_map[self.patient_id_column]
            row_idx = find_patient_row(
                ws,
                patient_id=ma_key,
                patient_col=patient_col,
                header_row=self.header_row,
            )
            if row_idx is None:
                row_idx = _next_data_row(ws, header_row=self.header_row)
                ws.cell(row=row_idx, column=patient_col, value=_coerce_patient_id(ma_key))

            stats = _apply_cell_updates(
                ws,
                row_idx=row_idx,
                header_map=header_map,
                updates=result.cell_updates,
                column_aliases=self.column_aliases,
                overwrite_policy=self.overwrite_policy,
                empty_sentinels=self.empty_sentinels,
            )
            wb.save(workbook_path)
        finally:
            wb.close()
        return stats


def _next_data_row(ws: Any, *, header_row: int) -> int:
    last = header_row
    for row_idx in range(header_row + 1, ws.max_row + 2):
        if any(
            ws.cell(row=row_idx, column=c).value is not None for c in range(1, ws.max_column + 1)
        ):
            last = row_idx
    return last + 1


def _coerce_patient_id(patient_id: str) -> str | int | float:
    try:
        if "." in patient_id:
            return float(patient_id)
        return int(patient_id)
    except ValueError:
        return patient_id


def _apply_cell_updates(
    ws: Any,
    *,
    row_idx: int,
    header_map: dict[str, int],
    updates: list[EcrfCellUpdate],
    column_aliases: dict[str, str],
    overwrite_policy: XlsOverwritePolicy,
    empty_sentinels: frozenset[str] | None,
) -> XlsWriteStats:
    stats = XlsWriteStats()
    for upd in updates:
        header = resolve_column_key(header_map, upd.column_key, column_aliases)
        if header is None:
            _LOG.debug("Colonne eCRF absente du gabarit : %s", upd.column_key)
            stats.skipped_missing_column += 1
            continue

        col_idx = header_map[header]
        live_value = ws.cell(row=row_idx, column=col_idx).value

        if not should_write_cell_value(
            live_value,
            policy=overwrite_policy,
            empty_sentinels=empty_sentinels,
        ):
            if overwrite_policy == XlsOverwritePolicy.EMPTY_ONLY:
                stats.skipped_filled += 1
            continue

        ws.cell(row=row_idx, column=col_idx, value=upd.value)
        stats.written += 1
    return stats


XlsExportPlaceholder = XlsEcrfExporter
