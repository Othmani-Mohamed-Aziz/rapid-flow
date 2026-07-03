"""Lecture gabarit eCRF et politique d'écrasement configurable."""

from __future__ import annotations

import logging
from enum import Enum
from pathlib import Path
from typing import Any

from app.etl.column_aliases import resolve_column_aliases
from app.etl.empty_sentinels import is_ecrf_cell_empty
from app.etl.xls_utils import build_header_map, find_patient_row, resolve_column_key

_LOG = logging.getLogger(__name__)


class XlsOverwritePolicy(str, Enum):
    EMPTY_ONLY = "empty_only"
    ALWAYS = "always"
    NEVER = "never"


class EcrfTemplateSnapshot:
    """Instantané lecture seule des valeurs eCRF d'un patient dans le gabarit MA."""

    def __init__(
        self,
        *,
        template_path: Path,
        ma_patient_key: str,
        sheet_name: str = "Global_CC",
        header_row: int = 2,
        patient_id_column: str = "ID_current_base",
        column_aliases: dict[str, str] | None = None,
        empty_sentinels: frozenset[str] | None = None,
    ) -> None:
        self.template_path = template_path
        self.ma_patient_key = ma_patient_key
        self.sheet_name = sheet_name
        self.header_row = header_row
        self.patient_id_column = patient_id_column
        self.column_aliases = column_aliases or resolve_column_aliases()
        self.empty_sentinels = empty_sentinels
        self._header_map: dict[str, int] = {}
        self.duplicate_headers: list[str] = []
        self._row_idx: int | None = None
        self._values: dict[str, Any] = {}
        self._loaded = False

    @classmethod
    def from_settings(cls, settings: Any, *, ma_patient_key: str) -> EcrfTemplateSnapshot:
        from app.etl.export import DEFAULT_ECRF_TEMPLATE_XLSX

        template = settings.export_xlsx_template_path
        return cls(
            template_path=Path(template) if template else DEFAULT_ECRF_TEMPLATE_XLSX,
            ma_patient_key=ma_patient_key,
            sheet_name=settings.export_xls_sheet_name,
            header_row=settings.export_xls_header_row,
            patient_id_column=settings.export_xls_patient_id_column,
            column_aliases=resolve_column_aliases(settings),
            empty_sentinels=settings.parsed_empty_sentinels(),
        )

    @property
    def available(self) -> bool:
        return self.template_path.is_file()

    def load(self) -> bool:
        if not self.available:
            return False
        if self._loaded:
            return True

        try:
            import openpyxl
        except ImportError:
            _LOG.warning("openpyxl absent : snapshot eCRF ignoré")
            return False

        wb = openpyxl.load_workbook(self.template_path, read_only=True, data_only=True)
        try:
            if self.sheet_name not in wb.sheetnames:
                _LOG.warning("Feuille %s absente du gabarit eCRF", self.sheet_name)
                return False
            ws = wb[self.sheet_name]
            header_result = build_header_map(ws, self.header_row)
            self._header_map = header_result.mapping
            self.duplicate_headers = header_result.duplicate_headers
            if self.patient_id_column not in self._header_map:
                return False
            patient_col = self._header_map[self.patient_id_column]
            self._row_idx = find_patient_row(
                ws,
                patient_id=self.ma_patient_key,
                patient_col=patient_col,
                header_row=self.header_row,
            )
            if self._row_idx is None:
                self._values = {}
            else:
                for header, col_idx in self._header_map.items():
                    self._values[header] = ws.cell(row=self._row_idx, column=col_idx).value
        finally:
            wb.close()

        self._loaded = True
        return True

    @property
    def patient_row_index(self) -> int | None:
        if not self._loaded:
            self.load()
        return self._row_idx

    def value_for_column(self, column_key: str) -> Any | None:
        if not self._loaded and not self.load():
            return None
        header = resolve_column_key(self._header_map, column_key, self.column_aliases)
        if header is None:
            return None
        return self._values.get(header)

    def is_column_filled(self, column_key: str) -> bool:
        return not is_ecrf_cell_empty(
            self.value_for_column(column_key),
            sentinels=self.empty_sentinels,
        )

    def filled_target_columns(self, target_columns: set[str] | frozenset[str]) -> set[str]:
        if not self.load():
            return set()
        return {col for col in target_columns if self.is_column_filled(col)}


def target_columns_for_document(study_schema: Any, doc_type: Any) -> set[str]:
    fields = study_schema.extractable_fields_for_document_type(doc_type)
    return {f.target_column for f in fields}
